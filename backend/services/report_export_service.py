from __future__ import annotations

import html
import re
import struct
import zlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from backend.services import report_service


SUPPORTED_FORMATS = {"markdown", "md", "html", "docx", "pdf"}
PDF_FONT_CANDIDATES = (
    "C:/Windows/Fonts/simsun.ttc",
    "C:/Windows/Fonts/simfang.ttf",
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/arialuni.ttf",
)

SECTION_NUMBERS = "一二三四五六七八九十"
FIELD_LABELS = {
    "ad_records": "投放记录",
    "avg_price": "均价",
    "clicks": "点击量",
    "comment_count": "评论数",
    "conversions": "成交量",
    "count": "数量",
    "ctr": "点击率",
    "customer_unit_price": "客单价",
    "forecast_avg_sales": "预测日均销量",
    "gmv": "销售额",
    "growth_score": "增长分",
    "impressions": "曝光量",
    "level": "优先级",
    "orders": "订单量",
    "positive_rate": "好评率",
    "price_gap": "价格差",
    "product_count": "商品数",
    "rating": "评分",
    "recent_avg_sales": "近期日均销量",
    "refund_count": "售后退款",
    "replenishment_count": "补货关注",
    "revenue": "成交金额",
    "risk_count": "风险商品",
    "risk_level": "风险等级",
    "risk_score": "风险分",
    "roi": "投产比",
    "sales_quantity": "销量",
    "slow_moving_count": "滞销关注",
    "spend": "投放花费",
    "stockout_count": "断货风险",
    "total_gmv": "销售额",
    "total_sales_quantity": "销量",
    "total_score": "综合分",
    "trend_direction": "趋势方向",
}
IDENTITY_KEYS = {
    "category",
    "content",
    "date",
    "desc",
    "differentiation_suggestion",
    "display",
    "insight",
    "keyword",
    "pain_point",
    "pricing_suggestion",
    "product_id",
    "product_name",
    "recommendation_actions",
    "title",
    "word",
}
MONEY_KEYS = {"avg_price", "competitor_price", "customer_unit_price", "gmv", "price_gap", "revenue", "spend", "total_gmv"}
PERCENT_KEYS = {"ctr", "positive_rate"}
METRIC_PRIORITY = [
    "total_gmv",
    "gmv",
    "revenue",
    "spend",
    "orders",
    "total_sales_quantity",
    "sales_quantity",
    "product_count",
    "total_score",
    "growth_score",
    "risk_score",
    "risk_count",
    "comment_count",
    "refund_count",
    "replenishment_count",
    "slow_moving_count",
    "stockout_count",
    "impressions",
    "clicks",
    "conversions",
    "ctr",
    "roi",
    "avg_price",
    "price_gap",
    "positive_rate",
    "rating",
    "count",
    "recent_avg_sales",
    "forecast_avg_sales",
    "risk_level",
    "trend_direction",
    "level",
]


@dataclass(frozen=True)
class ExportedReport:
    content: bytes
    mimetype: str
    filename: str
    format: str


class ReportExportError(Exception):
    pass


def export_report(report_id: str, export_format: str = "docx") -> ExportedReport | None:
    report = report_service.get_report(report_id)
    if not report:
        return None
    normalized = _normalize_format(export_format)
    if normalized == "markdown":
        return ExportedReport(
            content=_build_markdown(report).encode("utf-8-sig"),
            mimetype="text/markdown; charset=utf-8",
            filename=_filename(report, "md"),
            format=normalized,
        )
    if normalized == "html":
        return ExportedReport(
            content=_build_html(report).encode("utf-8"),
            mimetype="text/html; charset=utf-8",
            filename=_filename(report, "html"),
            format=normalized,
        )
    if normalized == "docx":
        return ExportedReport(
            content=_build_docx(report),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=_filename(report, "docx"),
            format=normalized,
        )
    if normalized == "pdf":
        return ExportedReport(
            content=_build_pdf(report),
            mimetype="application/pdf",
            filename=_filename(report, "pdf"),
            format=normalized,
        )
    raise ReportExportError(f"Unsupported report export format: {export_format}")


def _normalize_format(value: str | None) -> str:
    export_format = str(value or "docx").strip().lower()
    if export_format == "md":
        return "markdown"
    if export_format not in SUPPORTED_FORMATS:
        raise ReportExportError("format must be one of markdown, html, docx, or pdf")
    return export_format


def _build_markdown(report: dict[str, Any]) -> str:
    body = str(report.get("markdown") or "").strip()
    disclaimer = str(report.get("data_disclaimer") or "").strip()
    lines = [
        f"> 报告编号：{report.get('report_id', '--')}",
        f"> 生成时间：{report.get('generated_at', '--')}",
        f"> 报告范围：{_scope_text(report)}",
        "<!-- 店铺经营分析报告 -->",
        "",
        body,
    ]
    if disclaimer:
        lines.extend(["", "## 数据说明", disclaimer])
    return "\n".join(lines).strip() + "\n"


def _build_html(report: dict[str, Any]) -> str:
    title = str(report.get("title") or "ViralPick AI 经营分析报告")
    visuals = report.get("visuals") or {}
    sections = report.get("sections") or []
    kpis = visuals.get("kpis") or []
    trend = visuals.get("trend") or []
    categories = visuals.get("categories") or []
    actions = visuals.get("actions") or []
    max_trend = max([_float(item.get("value")) for item in trend] or [1]) or 1
    max_category = max([_float(item.get("value")) for item in categories] or [1]) or 1
    disclaimer = str(report.get("data_disclaimer") or "").strip()
    kpi_html = "\n".join(
        f"<article><span>{html.escape(str(item.get('label') or ''))}</span><strong>{html.escape(str(item.get('value') or ''))}</strong><em>{html.escape(str(item.get('desc') or ''))}</em></article>"
        for item in kpis
    )
    trend_html = "\n".join(
        f"<i style=\"height:{max(12, _float(item.get('value')) / max_trend * 100):.2f}%\"><b>{html.escape(str(item.get('display') or ''))}</b></i>"
        for item in trend
    )
    category_html = "\n".join(
        f"<div><span>{html.escape(str(item.get('category') or ''))}</span><strong>{html.escape(str(item.get('display') or ''))}</strong><i style=\"width:{max(5, _float(item.get('value')) / max_category * 100):.2f}%\"></i></div>"
        for item in categories
    )
    section_html = "\n".join(
        f"<article><span>{index + 1:02d}</span><h3>{html.escape(str(section.get('heading') or ''))}</h3><p>{html.escape(str(section.get('content') or ''))}</p></article>"
        for index, section in enumerate(sections)
    )
    action_html = "\n".join(
        f"<article><b>{index + 1}</b><div><strong>{html.escape(str(item.get('title') or ''))}</strong><span>{html.escape(str(item.get('content') or ''))}</span></div></article>"
        for index, item in enumerate(actions)
    )
    disclaimer_html = f"<section class=\"disclaimer\"><h3>数据说明</h3><p>{html.escape(disclaimer)}</p></section>" if disclaimer else ""
    return f"""<!doctype html>
<html lang="zh-CN">
  <head>
    <meta charset="utf-8" />
    <title>{html.escape(title)}</title>
    <style>
      :root {{ --ink:#102321; --muted:#5f716f; --line:#dbe8e5; --cyan:#00a99d; --blue:#4d7cff; --amber:#f0a51b; }}
      body {{ margin: 0; color: var(--ink); font: 14px/1.65 "Microsoft YaHei", "PingFang SC", sans-serif; background: #eef5f3; }}
      main {{ max-width: 1060px; margin: 32px auto; padding: 0 24px 44px; }}
      section {{ margin-top: 14px; }}
      .cover {{ display: grid; grid-template-columns: minmax(0,1fr) auto; gap: 16px; min-height: 142px; padding: 24px; color: #effffc; background: linear-gradient(135deg,#103f3a,#092724 58%,#172d44); border-radius: 12px; }}
      .cover span, .cover strong {{ color: #7ff5e6; font-weight: 900; }}
      .cover h1 {{ margin: 6px 0; font-size: 36px; line-height: 1.08; }}
      .cover p {{ margin: 0; max-width: 700px; color: #cbe5e1; }}
      .kpis, .visuals, .sections {{ display: grid; gap: 12px; }}
      .kpis {{ grid-template-columns: repeat(3,minmax(0,1fr)); }}
      .visuals {{ grid-template-columns: minmax(0,1.2fr) minmax(0,1fr); }}
      .sections {{ grid-template-columns: repeat(2,minmax(0,1fr)); }}
      article, .chart-card, .timeline, .disclaimer {{ padding: 16px; background: #fff; border: 1px solid var(--line); border-radius: 10px; box-shadow: 0 16px 40px rgba(28,55,52,.08); }}
      .kpis span, .kpis em, .timeline span, .disclaimer p {{ display: block; color: var(--muted); font-style: normal; }}
      .kpis strong {{ display:block; margin: 6px 0; font-size: 24px; }}
      h2, h3 {{ margin: 0 0 12px; color: var(--ink); line-height: 1.25; }}
      .mini-bars {{ height: 150px; display:flex; align-items:end; gap:8px; padding-top:26px; }}
      .mini-bars i {{ position:relative; flex:1 1 0; min-width:10px; background:linear-gradient(180deg,var(--cyan),#bff7ef); border-radius: 5px 5px 0 0; }}
      .mini-bars b {{ position:absolute; bottom:calc(100% + 5px); left:50%; transform:translateX(-50%); color:var(--muted); font-size:10px; white-space:nowrap; }}
      .category-bars {{ display:grid; gap:12px; }}
      .category-bars div {{ position:relative; display:grid; grid-template-columns:minmax(0,1fr) auto; gap:12px; padding-bottom:13px; }}
      .category-bars span {{ overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }}
      .category-bars strong {{ color:var(--amber); }}
      .category-bars i {{ position:absolute; left:0; bottom:0; height:6px; background:linear-gradient(90deg,var(--blue),var(--cyan)); border-radius:999px; }}
      .sections article > span {{ color: var(--amber); font-weight: 900; }}
      .sections p {{ margin:0; color:#314946; }}
      .timeline {{ display:grid; gap:10px; }}
      .timeline article {{ display:grid; grid-template-columns:34px minmax(0,1fr); gap:10px; padding:0; border:0; box-shadow:none; }}
      .timeline b {{ width:34px; height:34px; display:grid; place-items:center; color:#fff; background:var(--cyan); border-radius:50%; }}
      @media print {{ body {{ background:#fff; }} main {{ margin:0; max-width:none; }} article, .chart-card, .timeline, .disclaimer {{ box-shadow:none; }} }}
    </style>
  </head>
  <body>
    <main>
      <section class="cover">
        <div>
          <span>店铺经营分析报告</span>
          <h1>{html.escape(title)}</h1>
          <p>报告编号：{html.escape(str(report.get('report_id', '--')))} · 生成时间：{html.escape(str(report.get('generated_at', '--')))} · 报告范围：{html.escape(_scope_text(report))}</p>
        </div>
        <strong>{html.escape(str(report.get('scope_type') or 'global'))}</strong>
      </section>
      <section class="kpis">{kpi_html}</section>
      <section class="visuals">
        <article class="chart-card"><h2>收入走势</h2><div class="mini-bars">{trend_html}</div></article>
        <article class="chart-card"><h2>类目贡献</h2><div class="category-bars">{category_html}</div></article>
      </section>
      <section class="sections">{section_html}</section>
      <section class="timeline"><h2>执行安排</h2>{action_html}</section>
      {disclaimer_html}
    </main>
  </body>
</html>
"""


def _build_docx(report: dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise ReportExportError("python-docx is required for DOCX export") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    _configure_docx_styles(document, qn, OxmlElement, Pt, RGBColor)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(str(report.get("title") or "经营分析报告"))
    _format_docx_run(title_run, qn, Pt(18), bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("ViralPick AI 经营分析报告")
    _format_docx_run(subtitle_run, qn, Pt(10.5), bold=False)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "报告编号：{report_id}    生成时间：{generated_at}    报告范围：{scope}".format(
            report_id=report.get("report_id", "--"),
            generated_at=report.get("generated_at", "--"),
            scope=_scope_text(report),
        )
    )
    _format_docx_run(meta_run, qn, Pt(10.5), bold=False)
    document.add_paragraph()

    _append_structured_report_to_docx(
        document,
        report,
        qn,
        Pt,
        Cm,
        RGBColor,
        WD_CELL_VERTICAL_ALIGNMENT,
        WD_TABLE_ALIGNMENT,
        WD_ALIGN_PARAGRAPH,
    )
    disclaimer = str(report.get("data_disclaimer") or "").strip()
    if disclaimer:
        section_count = len(report.get("sections") or []) + 2
        _add_docx_heading(document, _section_title(section_count, "数据说明"), 1, qn, Pt)
        _add_docx_paragraph(document, disclaimer, qn, Pt)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf(report: dict[str, Any]) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ReportExportError("reportlab is required for PDF export") from exc

    font_name = _register_pdf_font(pdfmetrics, TTFont)
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("TitleCJK", parent=base["Title"], fontName=font_name, fontSize=20, leading=28, textColor=colors.black, alignment=1, spaceAfter=12),
        "meta": ParagraphStyle("MetaCJK", parent=base["BodyText"], fontName=font_name, fontSize=9.5, leading=15, textColor=colors.black, spaceAfter=12),
        "h1": ParagraphStyle("H1CJK", parent=base["Heading1"], fontName=font_name, fontSize=15, leading=21, textColor=colors.black, spaceBefore=12, spaceAfter=7),
        "h2": ParagraphStyle("H2CJK", parent=base["Heading2"], fontName=font_name, fontSize=12.5, leading=18, textColor=colors.black, spaceBefore=8, spaceAfter=5),
        "body": ParagraphStyle("BodyCJK", parent=base["BodyText"], fontName=font_name, fontSize=10.5, leading=17, textColor=colors.black, spaceAfter=5),
    }
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    story = [
        Paragraph(_pdf_escape(str(report.get("title") or "ViralPick AI 经营分析报告")), styles["title"]),
        Paragraph(
            "<br/>".join(
                [
                    f"报告编号：{_pdf_escape(str(report.get('report_id', '--')))}",
                    f"生成时间：{_pdf_escape(str(report.get('generated_at', '--')))}",
                    f"报告范围：{_pdf_escape(_scope_text(report))}",
                ]
            ),
            styles["meta"],
        ),
    ]
    _append_structured_report_to_pdf(story, report, styles, Paragraph, Spacer, Table, TableStyle, Image, colors, mm, font_name)
    disclaimer = str(report.get("data_disclaimer") or "").strip()
    if disclaimer:
        story.append(Spacer(1, 4 * mm))
        section_count = len(report.get("sections") or []) + 2
        story.append(Paragraph(_pdf_escape(_section_title(section_count, "数据说明")), styles["h1"]))
        story.append(Paragraph(_pdf_escape(disclaimer), styles["body"]))
    document.build(story)
    return buffer.getvalue()


def _append_markdown_to_docx(document, markdown: str) -> None:
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("> "):
            document.add_paragraph(line[2:].strip())
        elif line.startswith("<!--"):
            continue
        else:
            document.add_paragraph(line)


def _append_structured_report_to_docx(
    document,
    report: dict[str, Any],
    qn,
    pt_cls,
    cm_cls,
    rgb_cls,
    vertical_alignment_cls,
    table_alignment_cls,
    paragraph_alignment_cls,
) -> None:
    visuals = report.get("visuals") or {}
    sections = report.get("sections") or []
    kpis = visuals.get("kpis") or []

    _add_docx_heading(document, _section_title(1, "报告摘要"), 1, qn, pt_cls)
    _add_docx_paragraph(document, _summary_text(report), qn, pt_cls)
    if kpis:
        _add_docx_table(
            document,
            "核心指标",
            ["指标", "数值", "说明"],
            [[str(item.get("label") or ""), str(item.get("value") or ""), str(item.get("desc") or "")] for item in kpis],
            qn,
            pt_cls,
            rgb_cls,
            vertical_alignment_cls,
            table_alignment_cls,
        )
    _add_docx_chart(
        document,
        "收入趋势图",
        visuals.get("trend") or [],
        "trend",
        qn,
        pt_cls,
        cm_cls,
        paragraph_alignment_cls,
    )

    for index, section in enumerate(sections, 2):
        heading = str(section.get("heading") or f"分析项 {index - 1}")
        _add_docx_heading(document, _section_title(index, heading), 1, qn, pt_cls)
        _add_docx_paragraph(document, str(section.get("content") or ""), qn, pt_cls)
        evidence_rows = _evidence_rows(section.get("evidence"))
        if evidence_rows:
            _add_docx_table(
                document,
                "证据数据",
                ["项目", "指标", "数值"],
                evidence_rows,
                qn,
                pt_cls,
                rgb_cls,
                vertical_alignment_cls,
                table_alignment_cls,
            )
        chart_items, chart_type = _chart_items_for_section(section, visuals)
        chart_title = _chart_title_for_section(heading)
        _add_docx_chart(
            document,
            chart_title,
            chart_items,
            chart_type,
            qn,
            pt_cls,
            cm_cls,
            paragraph_alignment_cls,
        )


def _add_docx_table(
    document,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    qn,
    pt_cls,
    rgb_cls,
    vertical_alignment_cls,
    table_alignment_cls,
) -> None:
    if not rows:
        return
    _add_docx_heading(document, title, 2, qn, pt_cls)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = table_alignment_cls.CENTER
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(headers)]):
            cells[index].text = str(value)
    _format_docx_table(table, qn, pt_cls, rgb_cls, vertical_alignment_cls, header=True)


def _add_docx_chart(
    document,
    title: str,
    items: list[dict[str, Any]],
    chart_type: str,
    qn,
    pt_cls,
    cm_cls,
    paragraph_alignment_cls,
) -> None:
    image = _chart_png(items, chart_type, title)
    if not image:
        return
    _add_docx_heading(document, title, 2, qn, pt_cls)
    paragraph = document.add_paragraph()
    paragraph.alignment = paragraph_alignment_cls.CENTER
    run = paragraph.add_run()
    run.add_picture(BytesIO(image), width=cm_cls(15.6))


def _section_title(index: int, title: str) -> str:
    return f"{_chinese_number(index)}、{title}"


def _chinese_number(value: int) -> str:
    if 1 <= value <= 10:
        return SECTION_NUMBERS[value - 1]
    if 10 < value < 20:
        return "十" + SECTION_NUMBERS[value - 11]
    return str(value)


def _summary_text(report: dict[str, Any]) -> str:
    visuals = report.get("visuals") or {}
    kpis = {str(item.get("label") or ""): str(item.get("value") or "") for item in visuals.get("kpis") or []}
    trend = visuals.get("trend") or []
    categories = visuals.get("categories") or []
    leading_category = str((categories[0] or {}).get("category") or "重点类目") if categories else "重点类目"
    latest_trend = str((trend[-1] or {}).get("display") or "") if trend else ""
    trend_text = f"最近一期收入为 {latest_trend}，" if latest_trend else ""
    return (
        f"本报告范围为{_scope_text(report)}。截至生成时间，销售额 {kpis.get('销售额', '--')}，"
        f"订单量 {kpis.get('订单量', '--')}，投产比 {kpis.get('投产比', '--')}，"
        f"风险商品 {kpis.get('风险商品', '--')} 个，补货关注 {kpis.get('补货关注', '--')} 个。"
        f"{trend_text}{leading_category}贡献靠前，后续章节将按商品、增长、风险、口碑、竞品、投放、库存、趋势和执行动作逐项说明。"
    )


def _evidence_rows(evidence: Any, limit: int = 12) -> list[list[str]]:
    rows: list[list[str]] = []
    if isinstance(evidence, dict):
        for group, value in evidence.items():
            rows.extend(_rows_from_value(str(group), value))
    elif isinstance(evidence, list):
        rows.extend(_rows_from_value("数据项", evidence))
    else:
        rows.extend(_rows_from_value("数据项", evidence))
    return rows[:limit]


def _rows_from_value(group: str, value: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    if isinstance(value, list):
        for index, item in enumerate(value[:6], 1):
            if isinstance(item, dict):
                item_name = _item_name(item, group, index)
                metrics = _metric_pairs(item)
                for key, metric_value in metrics[:2]:
                    rows.append([item_name, _field_label(key), _display_metric(key, metric_value)])
            else:
                rows.append([_group_label(group), f"记录 {index}", str(item)])
        return rows
    if isinstance(value, dict):
        item_name = _item_name(value, group, 1)
        metrics = _metric_pairs(value)
        for key, metric_value in metrics[:8]:
            rows.append([item_name, _field_label(key), _display_metric(key, metric_value)])
        return rows
    if value not in (None, ""):
        rows.append([_group_label(group), "数值", str(value)])
    return rows


def _metric_pairs(item: dict[str, Any]) -> list[tuple[str, Any]]:
    pairs: list[tuple[str, Any]] = []
    for key in METRIC_PRIORITY:
        if key in item and _is_table_value(item.get(key)):
            pairs.append((key, item.get(key)))
    for key, value in item.items():
        if key in IDENTITY_KEYS or any(existing_key == key for existing_key, _ in pairs):
            continue
        if _is_table_value(value):
            pairs.append((key, value))
    return pairs


def _chart_items_for_section(section: dict[str, Any], visuals: dict[str, Any]) -> tuple[list[dict[str, Any]], str]:
    heading = str(section.get("heading") or "")
    evidence = section.get("evidence")
    if "经营概览" in heading:
        return _chart_items_from_kpis(visuals.get("kpis") or []), "categories"
    if "趋势" in heading:
        return visuals.get("trend") or [], "trend"
    if "商品" in heading:
        return _chart_items_from_records(_evidence_group(evidence, "top_products"), "product_name", ["total_gmv", "gmv", "total_sales_quantity"]), "categories"
    if "增长" in heading:
        return _chart_items_from_records(_evidence_group(evidence, "hot_scoring"), "product_name", ["total_score", "growth_score"]), "categories"
    if "风险" in heading:
        return _chart_items_from_records(_evidence_group(evidence, "risks"), "product_name", ["risk_score", "comment_count", "refund_count"]), "categories"
    if "口碑" in heading:
        positive = _chart_items_from_records(_evidence_group(evidence, "positive_keywords"), "word", ["count", "comment_count"])
        negative = _chart_items_from_records(_evidence_group(evidence, "negative_keywords"), "word", ["count", "comment_count"])
        return (positive + negative)[:10], "categories"
    if "竞品" in heading:
        return _chart_items_from_records(_evidence_group(evidence, "competitors"), "competitor_name", ["competitor_price", "price_gap", "sales_quantity"]), "categories"
    if "投放" in heading or "库存" in heading:
        return _chart_items_from_mapping(evidence), "categories"
    if "执行" in heading:
        return _chart_items_from_actions(_evidence_group(evidence, "today_actions")), "categories"
    return _chart_items_from_mapping(evidence), "categories"


def _chart_title_for_section(heading: str) -> str:
    if "经营概览" in heading:
        return "经营核心指标对比图"
    if "商品" in heading:
        return "商品销售额排行图"
    if "增长" in heading:
        return "增长机会评分图"
    if "风险" in heading:
        return "风险商品指标图"
    if "口碑" in heading:
        return "口碑关键词分布图"
    if "竞品" in heading:
        return "竞品指标对比图"
    if "投放" in heading:
        return "投放效果指标图"
    if "库存" in heading:
        return "库存售后指标图"
    if "趋势" in heading:
        return "未来销量趋势图"
    if "执行" in heading:
        return "执行优先级图"
    return "经营数据分析图"


def _chart_items_from_kpis(kpis: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items = []
    for item in kpis:
        value = _numeric_from_display(item.get("value"))
        if value > 0:
            items.append({"label": str(item.get("label") or ""), "value": value})
    return items[:8]


def _chart_items_from_records(records: Any, label_key: str, value_keys: list[str]) -> list[dict[str, Any]]:
    if not isinstance(records, list):
        return []
    items: list[dict[str, Any]] = []
    for index, record in enumerate(records[:8], 1):
        if not isinstance(record, dict):
            continue
        value = next((_float(record.get(key)) for key in value_keys if _float(record.get(key)) > 0), 0.0)
        if value <= 0:
            continue
        items.append({"label": str(record.get(label_key) or record.get("product_name") or record.get("title") or f"项目{index}"), "value": value})
    return items


def _chart_items_from_mapping(value: Any) -> list[dict[str, Any]]:
    mapping = value
    if isinstance(value, dict) and len(value) == 1:
        first = next(iter(value.values()))
        mapping = first if isinstance(first, dict) else value
    if not isinstance(mapping, dict):
        return []
    items: list[dict[str, Any]] = []
    for key in METRIC_PRIORITY:
        if key in mapping:
            numeric = _float(mapping.get(key))
            if numeric > 0:
                items.append({"label": _field_label(key), "value": numeric})
    for key, raw in mapping.items():
        if key in IDENTITY_KEYS or any(item["label"] == _field_label(key) for item in items):
            continue
        numeric = _float(raw)
        if numeric > 0:
            items.append({"label": _field_label(key), "value": numeric})
    return items[:8]


def _chart_items_from_actions(actions: Any) -> list[dict[str, Any]]:
    if not isinstance(actions, list):
        return []
    weights = {"high": 3, "normal": 2, "low": 1, "优先": 3, "跟进": 2}
    items = []
    for index, action in enumerate(actions[:6], 1):
        if not isinstance(action, dict):
            continue
        level = str(action.get("level") or "normal")
        items.append({"label": str(action.get("title") or f"动作{index}"), "value": weights.get(level, 2)})
    return items


def _evidence_group(evidence: Any, key: str) -> Any:
    if isinstance(evidence, dict):
        return evidence.get(key)
    return None


def _item_name(item: dict[str, Any], group: str, index: int) -> str:
    return str(
        item.get("product_name")
        or item.get("competitor_name")
        or item.get("title")
        or item.get("word")
        or item.get("keyword")
        or item.get("pain_point")
        or item.get("date")
        or item.get("category")
        or f"{_group_label(group)} {index}"
    )


def _group_label(group: str) -> str:
    return {
        "headline": "经营概览",
        "hot_scoring": "增长商品",
        "positive_keywords": "好评关键词",
        "negative_keywords": "差评关键词",
        "pain_points": "用户痛点",
        "risks": "风险商品",
        "top_products": "商品表现",
        "competitors": "竞品",
        "today_actions": "执行动作",
    }.get(group, group)


def _field_label(key: str) -> str:
    return FIELD_LABELS.get(key, key)


def _display_metric(key: str, value: Any) -> str:
    if key in MONEY_KEYS:
        return _currency(value)
    if key in PERCENT_KEYS:
        return _percent(value)
    if isinstance(value, (int, float)):
        return _compact_number(value)
    return str(value)


def _is_table_value(value: Any) -> bool:
    return isinstance(value, (str, int, float)) and value not in ("", None)


def _numeric_from_display(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value or ""))
    return _float(cleaned)


def _currency(value: Any) -> str:
    return f"¥{_float(value):,.0f}"


def _percent(value: Any) -> str:
    numeric = _float(value)
    if 0 <= numeric <= 1:
        numeric *= 100
    return f"{numeric:.1f}%"


def _compact_number(value: Any) -> str:
    numeric = _float(value)
    if abs(numeric - round(numeric)) < 0.0001:
        return f"{numeric:,.0f}"
    return f"{numeric:,.2f}"


def _configure_docx_styles(document, qn, oxml_element_cls, pt_cls, rgb_cls) -> None:
    for style_name, size, bold in [
        ("Normal", 10.5, False),
        ("Title", 18, True),
        ("Heading 1", 15, True),
        ("Heading 2", 12.5, True),
    ]:
        style = document.styles[style_name]
        style.font.name = "宋体"
        style.font.size = pt_cls(size)
        style.font.bold = bold
        style.font.color.rgb = rgb_cls(0, 0, 0)
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = oxml_element_cls("w:rFonts")
            r_pr.append(r_fonts)
        r_fonts.set(qn("w:eastAsia"), "宋体")
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def _format_docx_run(run, qn, size=None, bold: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.shared import RGBColor

    run.font.name = "宋体"
    if size is not None:
        run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_docx_heading(document, text: str, level: int, qn, pt_cls):
    paragraph = document.add_heading("", level=level)
    paragraph.paragraph_format.space_before = pt_cls(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = pt_cls(5)
    run = paragraph.add_run(text)
    _format_docx_run(run, qn, pt_cls(15 if level == 1 else 12.5), bold=True)
    return paragraph


def _add_docx_paragraph(document, text: str, qn, pt_cls):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = pt_cls(21)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = pt_cls(5)
    run = paragraph.add_run(text)
    _format_docx_run(run, qn, pt_cls(10.5), bold=False)
    return paragraph


def _format_docx_table(table, qn, pt_cls, rgb_cls, vertical_alignment_cls, header: bool = True) -> None:
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = vertical_alignment_cls.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = pt_cls(0)
                paragraph.paragraph_format.space_after = pt_cls(0)
                paragraph.paragraph_format.line_spacing = 1.25
                for run in paragraph.runs:
                    _format_docx_run(run, qn, pt_cls(10.5), bold=bool(header and row_index == 0))
                    run.font.color.rgb = rgb_cls(0, 0, 0)


def _append_markdown_to_pdf(story, markdown: str, styles, paragraph_cls, spacer_cls, mm_unit) -> None:
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or line.startswith("<!--"):
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            story.append(paragraph_cls(_pdf_escape(line[3:].strip()), styles["h1"]))
        elif line.startswith("### "):
            story.append(paragraph_cls(_pdf_escape(line[4:].strip()), styles["h2"]))
        elif line.startswith("- "):
            story.append(paragraph_cls(_pdf_escape(line[2:].strip()), styles["body"]))
        elif line.startswith("> "):
            story.append(paragraph_cls(_pdf_escape(line[2:].strip()), styles["meta"]))
        else:
            story.append(paragraph_cls(_pdf_escape(line), styles["body"]))
            story.append(spacer_cls(1, 1 * mm_unit))


def _append_structured_report_to_pdf(
    story,
    report: dict[str, Any],
    styles,
    paragraph_cls,
    spacer_cls,
    table_cls,
    table_style_cls,
    image_cls,
    colors_module,
    mm_unit,
    font_name: str,
) -> None:
    visuals = report.get("visuals") or {}
    sections = report.get("sections") or []
    kpis = visuals.get("kpis") or []

    story.append(paragraph_cls(_pdf_escape(_section_title(1, "报告摘要")), styles["h1"]))
    story.append(paragraph_cls(_pdf_escape(_summary_text(report)), styles["body"]))
    if kpis:
        data = [["指标", "数值", "经营含义"]] + [
            [str(item.get("label") or ""), str(item.get("value") or ""), str(item.get("desc") or "")]
            for item in kpis
        ]
        story.append(paragraph_cls("核心指标", styles["h2"]))
        story.append(_pdf_table(data, table_cls, table_style_cls, colors_module, font_name))
        story.append(spacer_cls(1, 3 * mm_unit))
    summary_chart = _pdf_chart(visuals.get("trend") or [], "trend", "收入趋势图", image_cls, mm_unit)
    if summary_chart:
        story.append(paragraph_cls("收入趋势图", styles["h2"]))
        story.append(summary_chart)
        story.append(spacer_cls(1, 3 * mm_unit))

    for index, section in enumerate(sections, 2):
        heading = str(section.get("heading") or f"分析项 {index - 1}")
        story.append(paragraph_cls(_pdf_escape(_section_title(index, heading)), styles["h1"]))
        story.append(paragraph_cls(_pdf_escape(str(section.get("content") or "")), styles["body"]))
        evidence_rows = _evidence_rows(section.get("evidence"))
        if evidence_rows:
            story.append(paragraph_cls("证据数据", styles["h2"]))
            story.append(_pdf_table([["项目", "指标", "数值"], *evidence_rows], table_cls, table_style_cls, colors_module, font_name))
            story.append(spacer_cls(1, 2 * mm_unit))
        chart_items, chart_type = _chart_items_for_section(section, visuals)
        chart_title = _chart_title_for_section(heading)
        chart = _pdf_chart(chart_items, chart_type, chart_title, image_cls, mm_unit)
        if chart:
            story.append(paragraph_cls(chart_title, styles["h2"]))
            story.append(chart)
            story.append(spacer_cls(1, 3 * mm_unit))


def _pdf_chart(items: list[dict[str, Any]], chart_type: str, title: str, image_cls, mm_unit):
    image = _chart_png(items, chart_type, title)
    if not image:
        return None
    return image_cls(BytesIO(image), width=160 * mm_unit, height=55 * mm_unit)


def _pdf_table(data, table_cls, table_style_cls, colors_module, font_name: str):
    table = table_cls(data, hAlign="LEFT")
    table.setStyle(
        table_style_cls(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors_module.HexColor("#f2f2f2")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors_module.black),
                ("GRID", (0, 0), (-1, -1), 0.4, colors_module.black),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("LEADING", (0, 0), (-1, -1), 13),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors_module.white, colors_module.HexColor("#fafafa")]),
            ]
        )
    )
    return table


def _chart_png(items: list[dict[str, Any]], chart_type: str, title: str = "经营数据分析图") -> bytes | None:
    cleaned = [
        {**item, "value": _float(item.get("value")), "label": _chart_label(item)}
        for item in items
        if _float(item.get("value")) >= 0
    ]
    values = [float(item.get("value") or 0) for item in cleaned]
    if not values:
        return None
    rendered = _matplotlib_chart_png(cleaned, chart_type, title)
    if rendered:
        return rendered
    width, height = 760, 260
    canvas = bytearray([246, 251, 250] * width * height)

    def rect(x1: int, y1: int, x2: int, y2: int, color: tuple[int, int, int]) -> None:
        x1 = max(0, min(width, x1))
        x2 = max(0, min(width, x2))
        y1 = max(0, min(height, y1))
        y2 = max(0, min(height, y2))
        if x2 <= x1 or y2 <= y1:
            return
        for y in range(y1, y2):
            row = (y * width + x1) * 3
            for _ in range(x1, x2):
                canvas[row : row + 3] = bytes(color)
                row += 3

    def dot(x: int, y: int, color: tuple[int, int, int]) -> None:
        rect(x - 3, y - 3, x + 4, y + 4, color)

    def line(x1: int, y1: int, x2: int, y2: int, color: tuple[int, int, int]) -> None:
        dx = abs(x2 - x1)
        dy = -abs(y2 - y1)
        sx = 1 if x1 < x2 else -1
        sy = 1 if y1 < y2 else -1
        error = dx + dy
        x, y = x1, y1
        while True:
            rect(x - 1, y - 1, x + 2, y + 2, color)
            if x == x2 and y == y2:
                break
            next_error = 2 * error
            if next_error >= dy:
                error += dy
                x += sx
            if next_error <= dx:
                error += dx
                y += sy

    rect(0, 0, width, height, (246, 251, 250))
    rect(0, 0, width, 6, (0, 169, 157))
    rect(28, 28, width - 28, height - 28, (255, 255, 255))
    max_value = max(values) or 1

    for index in range(5):
        y = 42 + index * 38
        rect(40, y, width - 40, y + 1, (220, 232, 229))

    if chart_type == "categories":
        visible = cleaned[:6]
        bar_gap = 18
        bar_height = 18
        start_y = 54
        for index, item in enumerate(visible):
            value = float(item.get("value") or 0)
            bar_width = int((width - 160) * (value / max_value))
            y = start_y + index * (bar_height + bar_gap)
            rect(90, y, 90 + max(16, bar_width), y + bar_height, (0, 214, 190))
            rect(90 + max(16, bar_width), y, 92 + max(16, bar_width), y + bar_height, (255, 209, 102))
    else:
        visible = cleaned[-12:]
        gap = 10
        bar_area_width = width - 100
        bar_width = max(10, int((bar_area_width - gap * (len(visible) - 1)) / max(len(visible), 1)))
        base_y = height - 42
        points: list[tuple[int, int]] = []
        for index, item in enumerate(visible):
            value = float(item.get("value") or 0)
            bar_height = int(158 * (value / max_value))
            x = 52 + index * (bar_width + gap)
            rect(x, base_y - max(8, bar_height), x + bar_width, base_y, (0, 214, 190))
            rect(x, base_y - max(8, bar_height), x + bar_width, base_y - max(8, bar_height) + 4, (255, 209, 102))
            points.append((x + bar_width // 2, base_y - max(8, bar_height)))
        if chart_type == "trend" and len(points) > 1:
            for index in range(1, len(points)):
                line(points[index - 1][0], points[index - 1][1], points[index][0], points[index][1], (91, 140, 255))
            for x, y in points:
                dot(x, y, (255, 209, 102))

    return _png_bytes(width, height, canvas)


def _matplotlib_chart_png(items: list[dict[str, Any]], chart_type: str, title: str) -> bytes | None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.font_manager import FontProperties
        from matplotlib.ticker import FuncFormatter
    except Exception:
        return None

    visible = items[-12:] if chart_type == "trend" else items[:8]
    if not visible:
        return None
    labels = [_short_label(str(item.get("label") or f"项目{index + 1}")) for index, item in enumerate(visible)]
    values = [float(item.get("value") or 0) for item in visible]
    font = _chart_font(FontProperties)
    axis_label = _chart_axis_label(title)
    buffer = BytesIO()

    try:
        plt.rcParams["axes.unicode_minus"] = False
        fig, ax = plt.subplots(figsize=(8.4, 3.2), dpi=160)
        fig.patch.set_facecolor("#ffffff")
        ax.set_facecolor("#ffffff")
        ax.grid(axis="y" if chart_type == "trend" else "x", color="#d9e4e1", linewidth=0.8, alpha=0.9)
        ax.set_axisbelow(True)
        for spine in ["top", "right"]:
            ax.spines[spine].set_visible(False)
        ax.spines["left"].set_color("#5f716f")
        ax.spines["bottom"].set_color("#5f716f")
        ax.tick_params(colors="#304744", labelsize=8)
        ax.yaxis.set_major_formatter(FuncFormatter(lambda value, _pos: _format_axis_value(value)))
        ax.xaxis.set_major_formatter(FuncFormatter(lambda value, _pos: _format_axis_value(value) if chart_type != "trend" else str(int(value)) if value == int(value) else ""))

        if chart_type == "categories":
            ordered = list(zip(labels, values))
            ordered.reverse()
            y_labels = [item[0] for item in ordered]
            y_values = [item[1] for item in ordered]
            y_positions = list(range(len(y_labels)))
            bars = ax.barh(y_positions, y_values, color="#00a99d", edgecolor="#087c75", linewidth=0.8)
            if y_values:
                ax.set_xlim(0, max(y_values) * 1.22)
            ax.set_xlabel(axis_label, fontproperties=font, fontsize=9, color="#304744")
            ax.set_ylabel("项目名称", fontproperties=font, fontsize=9, color="#304744")
            ax.set_yticks(y_positions)
            ax.set_yticklabels(y_labels, fontproperties=font, fontsize=8)
            for bar, value in zip(bars, y_values):
                ax.text(
                    bar.get_width(),
                    bar.get_y() + bar.get_height() / 2,
                    "  " + _format_axis_value(value),
                    va="center",
                    ha="left",
                    fontsize=8,
                    color="#102321",
                    fontproperties=font,
                )
        else:
            x_positions = list(range(len(values)))
            bars = ax.bar(x_positions, values, color="#00a99d", edgecolor="#087c75", linewidth=0.8, label="指标值")
            ax.plot(x_positions, values, color="#315cff", marker="o", linewidth=2, markersize=4, label="趋势线")
            if values:
                ax.set_ylim(0, max(values) * 1.24)
            ax.set_xticks(x_positions)
            ax.set_xticklabels(labels, rotation=35, ha="right", fontproperties=font, fontsize=8)
            ax.set_xlabel("日期/周期", fontproperties=font, fontsize=9, color="#304744")
            ax.set_ylabel(axis_label, fontproperties=font, fontsize=9, color="#304744")
            ax.legend(prop=font, frameon=False, fontsize=8, loc="upper left")
            for bar, value in zip(bars, values):
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    bar.get_height(),
                    _format_axis_value(value),
                    va="bottom",
                    ha="center",
                    fontsize=7,
                    color="#102321",
                    fontproperties=font,
                )

        ax.set_title(title, fontproperties=font, fontsize=13, fontweight="bold", color="#102321", pad=10)
        for label in ax.get_xticklabels() + ax.get_yticklabels():
            label.set_fontproperties(font)
            label.set_fontsize(8)
        fig.tight_layout(pad=1.2)
        fig.savefig(buffer, format="png", facecolor="#ffffff", bbox_inches="tight")
        plt.close(fig)
        return buffer.getvalue()
    except Exception:
        try:
            plt.close("all")
        except Exception:
            pass
        return None


def _chart_font(font_properties_cls):
    for candidate in PDF_FONT_CANDIDATES:
        if Path(candidate).exists():
            return font_properties_cls(fname=candidate)
    return font_properties_cls()


def _chart_label(item: dict[str, Any]) -> str:
    return str(
        item.get("label")
        or item.get("name")
        or item.get("date")
        or item.get("category")
        or item.get("product_name")
        or item.get("title")
        or item.get("word")
        or "项目"
    )


def _chart_axis_label(title: str) -> str:
    if re.search(r"销售额|收入|金额|投放|花费|GMV|gmv", title):
        return "金额（元）"
    if re.search(r"销量|订单|评论|库存|售后|曝光|点击|成交|数量", title):
        return "数量"
    if re.search(r"评分|分值|优先级|投产比|好评率|点击率|ROI", title):
        return "分值/比例"
    return "数值"


def _short_label(value: str, max_len: int = 12) -> str:
    text = str(value or "").strip()
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def _format_axis_value(value: Any) -> str:
    numeric = _float(value)
    absolute = abs(numeric)
    if absolute >= 100_000_000:
        return f"{numeric / 100_000_000:.1f}亿"
    if absolute >= 10_000:
        return f"{numeric / 10_000:.1f}万"
    if absolute >= 1000:
        return f"{numeric:,.0f}"
    if absolute >= 100:
        return f"{numeric:.0f}"
    if absolute >= 10:
        return f"{numeric:.1f}"
    return f"{numeric:.2f}".rstrip("0").rstrip(".")


def _png_bytes(width: int, height: int, rgb: bytearray) -> bytes:
    def chunk(chunk_type: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + chunk_type
            + data
            + struct.pack(">I", zlib.crc32(chunk_type + data) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + bytes(rgb[row * width * 3 : (row + 1) * width * 3]) for row in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def _markdown_to_html(markdown: str) -> str:
    lines: list[str] = []
    in_list = False
    for raw in markdown.splitlines():
        line = raw.strip()
        if line.startswith("<!--"):
            continue
        if not line:
            if in_list:
                lines.append("</ul>")
                in_list = False
            continue
        if line.startswith("- "):
            if not in_list:
                lines.append("<ul>")
                in_list = True
            lines.append(f"<li>{html.escape(line[2:].strip())}</li>")
            continue
        if in_list:
            lines.append("</ul>")
            in_list = False
        if line.startswith("# "):
            lines.append(f"<h1>{html.escape(line[2:].strip())}</h1>")
        elif line.startswith("## "):
            lines.append(f"<h2>{html.escape(line[3:].strip())}</h2>")
        elif line.startswith("### "):
            lines.append(f"<h3>{html.escape(line[4:].strip())}</h3>")
        elif line.startswith("> "):
            lines.append(f"<blockquote>{html.escape(line[2:].strip())}</blockquote>")
        else:
            lines.append(f"<p>{html.escape(line)}</p>")
    if in_list:
        lines.append("</ul>")
    return "\n".join(lines)


def _register_pdf_font(pdfmetrics, ttfont_cls) -> str:
    font_name = "ViralPickCJK"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    for candidate in PDF_FONT_CANDIDATES:
        if Path(candidate).exists():
            try:
                pdfmetrics.registerFont(ttfont_cls(font_name, candidate))
                return font_name
            except Exception:
                continue
    return "Helvetica"


def _pdf_escape(value: str) -> str:
    return html.escape(value, quote=False).replace("\n", "<br/>")


def _float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _scope_text(report: dict[str, Any]) -> str:
    if report.get("product_id"):
        return f"单品：{report.get('product_id')}"
    if report.get("category"):
        return f"品类：{report.get('category')}"
    return "全店"


def _filename(report: dict[str, Any], extension: str) -> str:
    title = str(report.get("title") or "ViralPick-AI-report")
    safe_title = re.sub(r'[\\/:*?"<>|\r\n]+', "_", title).strip(" ._")
    return f"{(safe_title[:80] or 'ViralPick-AI-report')}-{report.get('report_id', 'report')}.{extension}"
