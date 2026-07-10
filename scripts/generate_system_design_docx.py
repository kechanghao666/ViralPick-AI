from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
SOURCE = BASE_DIR / "docs" / "system_design.md"
OUTPUT = BASE_DIR / "docs" / "ViralPick AI 系统设计文档.docx"
DESIGN_DIR = BASE_DIR / "docs" / "design"

FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")
BOLD_FONT_PATH = Path(r"C:\Windows\Fonts\msyhbd.ttc")

FONT_CN = "微软雅黑"
COLOR_DARK = "0B1F2A"
COLOR_TEAL = "00D6C9"
COLOR_TEAL_DARK = "0F766E"
COLOR_BLUE = "2563EB"
COLOR_BG = "F6FBFA"
COLOR_CARD = "FFFFFF"
COLOR_LINE = "B8D7D3"
COLOR_MUTED = "64748B"
COLOR_TABLE_HEAD = "0F766E"
COLOR_TABLE_ALT = "F3FAF9"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = BOLD_FONT_PATH if bold and BOLD_FONT_PATH.exists() else FONT_PATH
    return ImageFont.truetype(str(path), size)


def draw_round(draw: ImageDraw.ImageDraw, xy, fill, outline=COLOR_LINE, width=2, radius=22) -> None:
    fill_color = fill if str(fill).startswith("#") else "#" + str(fill)
    outline_color = outline if str(outline).startswith("#") else "#" + str(outline)
    draw.rounded_rectangle(xy, radius=radius, fill=fill_color, outline=outline_color, width=width)


def draw_center_text(draw: ImageDraw.ImageDraw, xy, text: str, size=26, fill="#0B1F2A", bold=False) -> None:
    x1, y1, x2, y2 = xy
    f = font(size, bold)
    lines = text.split("\n")
    total_h = len(lines) * (size + 8) - 8
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=f)
        x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, fill=fill, font=f)
        y += size + 8


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color="#0F766E", width=4) -> None:
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        head = [(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)]
    else:
        head = [(x2, y2), (x2 + 18, y2 - 10), (x2 + 18, y2 + 10)]
    draw.polygon(head, fill=color)


def draw_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((60, 42), title, fill="#083B3A", font=font(40, True))
    draw.text((62, 95), subtitle, fill="#475569", font=font(22))
    draw.line((60, 128, 1540, 128), fill="#00D6C9", width=5)


def generate_architecture_image() -> Path:
    path = DESIGN_DIR / "system_architecture.png"
    img = Image.new("RGB", (1600, 980), "#" + COLOR_BG)
    draw = ImageDraw.Draw(img)
    draw_title(draw, "技术架构图", "前端展示、Flask API、业务服务、仓储层与 SQLite 数据层分工")

    boxes = [
        ((70, 210, 310, 370), "用户浏览器\n经营人员"),
        ((390, 190, 660, 390), "前端表现层\nHTML / CSS / JS\nECharts"),
        ((740, 190, 1010, 390), "API 路由层\nFlask Blueprint\n统一响应"),
        ((1090, 150, 1500, 430), "业务服务层\n导入 / 分析 / 评分\n预测 / 报告 / 助手"),
        ((740, 560, 1010, 750), "仓储层\nrepositories\nSQL 查询封装"),
        ((1090, 560, 1500, 750), "数据层\nSQLite 数据库\nCSV / Excel 文件"),
    ]
    for xy, text in boxes:
        draw_round(draw, xy, COLOR_CARD)
        draw_center_text(draw, xy, text, 25, "#0B1F2A", True)

    draw_arrow(draw, (310, 290), (390, 290))
    draw_arrow(draw, (660, 290), (740, 290))
    draw_arrow(draw, (1010, 290), (1090, 290))
    draw_arrow(draw, (1300, 430), (1300, 560))
    draw_arrow(draw, (1090, 655), (1010, 655))
    draw_arrow(draw, (875, 560), (875, 390))

    modules = [
        "数据接入中心", "经营驾驶舱", "商品增长中心", "用户口碑中心",
        "投放内容中心", "趋势预测中心", "经营报告中心", "经营助手"
    ]
    x, y = 110, 820
    for item in modules:
        draw_round(draw, (x, y, x + 250, y + 72), "E6FFFB", COLOR_TEAL_DARK, 2, 16)
        draw_center_text(draw, (x, y, x + 250, y + 72), item, 21, "#064E4A", True)
        x += 180 if item == "经营助手" else 180
        if x + 250 > 1540:
            x = 110
            y += 84
    img.save(path)
    return path


def generate_er_image() -> Path:
    path = DESIGN_DIR / "database_er.png"
    img = Image.new("RGB", (1700, 1100), "#" + COLOR_BG)
    draw = ImageDraw.Draw(img)
    draw_title(draw, "数据库 ER 图", "以 products 为核心，连接销售、订单、口碑、库存、投放、预测和报告数据")

    center = (690, 430, 1010, 610)
    draw_round(draw, center, "DBEAFE", COLOR_BLUE, 3, 24)
    draw_center_text(draw, center, "products\n商品主表", 30, "#1E3A8A", True)

    tables = [
        ((80, 190, 360, 310), "sales_daily\n每日销售"),
        ((420, 190, 650, 310), "comments\n用户评论"),
        ((1050, 190, 1280, 310), "orders_daily\n订单经营"),
        ((1340, 190, 1620, 310), "inventory_daily\n库存"),
        ((80, 430, 360, 550), "after_sales\n售后"),
        ((420, 700, 650, 820), "ad_metrics\n广告投放"),
        ((720, 760, 980, 880), "traffic_metrics\n流量转化"),
        ((1050, 700, 1320, 820), "content_metrics\n内容运营"),
        ((1340, 430, 1620, 550), "forecast_results\n预测结果"),
        ((80, 700, 360, 820), "competitors\n市场对照"),
        ((1340, 700, 1620, 820), "analysis_reports\n经营报告"),
    ]
    for xy, text in tables:
        draw_round(draw, xy, COLOR_CARD)
        draw_center_text(draw, xy, text, 24, "#0B1F2A", True)
        sx = (xy[0] + xy[2]) // 2
        sy = (xy[1] + xy[3]) // 2
        cx = (center[0] + center[2]) // 2
        cy = (center[1] + center[3]) // 2
        draw.line((sx, sy, cx, cy), fill="#0F766E", width=3)

    support = [
        ((250, 940, 510, 1020), "import_previews"),
        ((570, 940, 830, 1020), "import_jobs"),
        ((890, 940, 1150, 1020), "scoring_models"),
        ((1210, 940, 1470, 1020), "report_shares / schedules"),
    ]
    for xy, text in support:
        draw_round(draw, xy, "F8FAFC", COLOR_LINE, 2, 16)
        draw_center_text(draw, xy, text, 22, "#334155", True)
    draw_arrow(draw, (510, 980), (570, 980), "#64748B", 3)
    img.save(path)
    return path


def generate_ui_image() -> Path:
    path = DESIGN_DIR / "ui_prototype.png"
    img = Image.new("RGB", (1700, 1050), "#071A1B")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1700, 1050), fill="#071A1B")
    draw.text((60, 42), "界面原型图", fill="#EAF7F5", font=font(42, True))
    draw.text((62, 96), "左侧中心导航 + 顶部搜索操作 + 主业务大屏 + 全局经营助手", fill="#94A3B8", font=font(23))
    draw.line((60, 130, 1640, 130), fill="#00D6C9", width=5)

    draw_round(draw, (60, 180, 350, 970), "062224", "134E4A", 2, 18)
    draw.text((100, 220), "ViralPick AI", fill="#EAF7F5", font=font(28, True))
    nav = ["经营驾驶舱", "数据接入中心", "商品增长中心", "用户口碑中心", "投放内容中心", "趋势预测中心", "经营报告中心"]
    y = 300
    for idx, item in enumerate(nav, 1):
        fill = "0F4F4B" if idx == 1 else "082E30"
        draw_round(draw, (95, y, 315, y + 58), fill, "1B6B66", 2, 12)
        draw.text((115, y + 15), f"{idx:02d}", fill="#00D6C9", font=font(16, True))
        draw.text((165, y + 13), item, fill="#EAF7F5", font=font(19, True))
        y += 80

    draw_round(draw, (390, 180, 1620, 260), "0A2228", "1D4F55", 2, 16)
    draw.text((430, 205), "经营驾驶舱", fill="#EAF7F5", font=font(28, True))
    draw_round(draw, (1180, 202, 1500, 242), "071A1B", "1D4F55", 2, 10)
    draw.text((1210, 212), "搜索商品、类目、店铺", fill="#94A3B8", font=font(17))
    draw_round(draw, (1510, 202, 1600, 242), "00D6C9", "00D6C9", 2, 10)
    draw_center_text(draw, (1510, 202, 1600, 242), "生成汇报", 17, "#063B3A", True)

    draw_round(draw, (390, 290, 1160, 500), "092B2E", "1D4F55", 2, 18)
    draw.text((430, 325), "今日经营概况", fill="#EAF7F5", font=font(36, True))
    for i, label in enumerate(["销售额", "订单量", "客单价", "投产比"]):
        x = 430 + i * 175
        draw_round(draw, (x, 400, x + 150, 465), "13363A", "1D4F55", 2, 10)
        draw.text((x + 18, 414), label, fill="#94A3B8", font=font(16))
        draw.text((x + 18, 438), "核心指标", fill="#EAF7F5", font=font(18, True))
    draw_round(draw, (1190, 290, 1620, 500), "09212B", "1D4F55", 2, 18)
    draw_center_text(draw, (1190, 290, 1620, 500), "经营信号\n风险与动作", 30, "#EAF7F5", True)

    panels = [
        (390, 530, 1040, 820, "近 30 日销售趋势"),
        (1070, 530, 1620, 820, "今日动作 / 风险预警"),
        (390, 850, 770, 970, "增长机会"),
        (800, 850, 1180, 970, "类目贡献"),
        (1210, 850, 1620, 970, "库存售后"),
    ]
    for x1, y1, x2, y2, title in panels:
        draw_round(draw, (x1, y1, x2, y2), "092B2E", "1D4F55", 2, 16)
        draw.text((x1 + 24, y1 + 20), title, fill="#EAF7F5", font=font(21, True))
    draw_round(draw, (1430, 925, 1640, 1000), "9DF273", "9DF273", 2, 38)
    draw_center_text(draw, (1430, 925, 1640, 1000), "经营助手", 22, "#063B3A", True)
    img.save(path)
    return path


def generate_api_flow_image() -> Path:
    path = DESIGN_DIR / "api_flow.png"
    img = Image.new("RGB", (1600, 950), "#" + COLOR_BG)
    draw = ImageDraw.Draw(img)
    draw_title(draw, "接口与业务流程图", "前端页面通过统一 API 调用后端服务，服务层完成分析并写入 SQLite")

    lanes = [
        ((80, 190, 380, 760), "前端页面"),
        ((460, 190, 760, 760), "Flask API"),
        ((840, 190, 1140, 760), "服务层"),
        ((1220, 190, 1520, 760), "SQLite / 文件"),
    ]
    for xy, text in lanes:
        draw_round(draw, xy, "FFFFFF", COLOR_LINE, 2, 18)
        draw.text((xy[0] + 70, xy[1] + 28), text, fill="#0B1F2A", font=font(28, True))

    flows = [
        (250, 300, "上传文件", "POST /imports/preview", "字段识别", "预览缓存"),
        (250, 430, "确认导入", "POST /imports/commit", "清洗写库", "业务表"),
        (250, 560, "查看看板", "GET /dashboard/executive", "经营汇总", "读取指标"),
        (250, 690, "生成报告", "POST /reports/generate", "报告编排", "保存报告"),
    ]
    for y, y_api, a, b, c, d in flows:
        draw.text((115, y - 18), a, fill="#0B1F2A", font=font(21, True))
        draw.text((500, y_api - 18), b, fill="#0F766E", font=font(18, True))
        draw.text((900, y_api - 18), c, fill="#0B1F2A", font=font(21, True))
        draw.text((1305, y_api - 18), d, fill="#0B1F2A", font=font(21, True))
        draw_arrow(draw, (330, y), (470, y_api), "#0F766E", 4)
        draw_arrow(draw, (720, y_api), (850, y_api), "#0F766E", 4)
        draw_arrow(draw, (1110, y_api), (1230, y_api), "#0F766E", 4)
    img.save(path)
    return path


def set_run_font(run, size: int | None = None, bold: bool = False, color: str | None = None) -> None:
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "C7DAD8", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_border(paragraph, color: str = COLOR_TEAL, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            if i == 0:
                set_cell_shading(cell, COLOR_TABLE_HEAD)
            elif i % 2 == 0:
                set_cell_shading(cell, COLOR_TABLE_ALT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value.strip())
            set_run_font(run, 9, i == 0, "FFFFFF" if i == 0 else COLOR_DARK)
    doc.add_paragraph()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = FONT_CN
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 18, COLOR_TEAL_DARK), ("Heading 2", 14, COLOR_DARK), ("Heading 3", 12, COLOR_TEAL_DARK)]:
        style = styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("ViralPick AI 系统设计文档")
    set_run_font(run, 9, False, COLOR_MUTED)
    return doc


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    run = p.add_run("ViralPick AI")
    set_run_font(run, 34, True, COLOR_TEAL_DARK)
    p = doc.add_paragraph()
    run = p.add_run("系统设计文档")
    set_run_font(run, 25, True, COLOR_DARK)
    set_paragraph_border(p, COLOR_TEAL, "24")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("数据库设计 · ER 图 · 界面原型图 · 技术架构图 · 接口设计")
    set_run_font(run, 14, True, COLOR_DARK)
    add_table(doc, [
        ["项目", "内容"],
        ["中文定位", "基于 Python 的多平台电商智能经营与增长决策平台"],
        ["技术路线", "Flask + SQLite + pandas/numpy + scikit-learn + ECharts"],
        ["文档范围", "数据库设计、界面原型、技术架构、模块分层、接口设计"],
        ["生成日期", date.today().isoformat()],
    ])
    doc.add_page_break()


def parse_markdown(md: str) -> tuple[list, list[str]]:
    elements: list = []
    toc: list[str] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith(">"):
            quotes = []
            while i < len(lines) and lines[i].startswith(">"):
                quotes.append(lines[i].lstrip("> ").strip())
                i += 1
            elements.append(("quote", quotes))
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", line)
            if match:
                elements.append(("image", match.group(1), match.group(2)))
            i += 1
            continue
        if line.startswith("```"):
            lang = line.strip("`").strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            elements.append(("code", lang, block))
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            if level == 2:
                toc.append(text)
            elements.append(("heading", level, text))
            i += 1
            continue
        if line.startswith("|"):
            raw_rows = []
            while i < len(lines) and lines[i].startswith("|"):
                raw_rows.append(lines[i])
                i += 1
            rows = []
            for raw in raw_rows:
                cells = [cell.strip() for cell in raw.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
                    continue
                rows.append(cells)
            elements.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            elements.append(("bullets", items))
            continue
        paragraph = [line]
        i += 1
        while i < len(lines):
            current = lines[i].rstrip()
            if not current.strip() or current.startswith(("#", "|", "- ", ">", "```", "![")):
                break
            paragraph.append(current)
            i += 1
        elements.append(("paragraph", " ".join(x.strip() for x in paragraph)))
    return elements, toc


def add_text_run(paragraph, text: str, size: int = 10.5, bold: bool = False) -> None:
    text = text.replace("`", "").replace("**", "")
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, COLOR_DARK)


def render_doc(doc: Document, elements: list, toc: list[str]) -> None:
    p = doc.add_paragraph("目录", style="Heading 1")
    set_paragraph_border(p)
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        add_text_run(p, item, 10.5)
    doc.add_page_break()

    skip_title = True
    for element in elements:
        kind = element[0]
        if kind == "heading":
            _, level, text = element
            if level == 1 and skip_title:
                skip_title = False
                continue
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            p = doc.add_paragraph(text, style=style)
            if level == 2:
                set_paragraph_border(p)
            continue
        if kind == "quote":
            for line in element[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                add_text_run(p, line, 9.5, True)
            continue
        if kind == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(5)
            add_text_run(p, element[1], 10.5)
            continue
        if kind == "bullets":
            for item in element[1]:
                p = doc.add_paragraph(style="List Bullet")
                add_text_run(p, item, 10)
            continue
        if kind == "table":
            add_table(doc, element[1])
            continue
        if kind == "code":
            lang, block = element[1], element[2]
            if lang in {"mermaid", "text", "json"}:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run("图示说明见源文档对应代码块或下方图片。")
                set_run_font(run, 9, False, COLOR_MUTED)
            continue
        if kind == "image":
            alt, rel = element[1], element[2]
            image_path = BASE_DIR / "docs" / rel
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(alt)
                set_run_font(run, 10, True, COLOR_TEAL_DARK)
                doc.add_picture(str(image_path), width=Cm(16.5))
            continue


def main() -> None:
    DESIGN_DIR.mkdir(parents=True, exist_ok=True)
    generate_architecture_image()
    generate_er_image()
    generate_ui_image()
    generate_api_flow_image()
    md = SOURCE.read_text(encoding="utf-8")
    elements, toc = parse_markdown(md)
    doc = setup_document()
    add_cover(doc)
    render_doc(doc, elements, toc)
    doc.core_properties.title = "ViralPick AI 系统设计文档"
    doc.core_properties.subject = "数据库设计、ER 图、界面原型、技术架构和接口设计"
    doc.core_properties.author = "ViralPick AI"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
