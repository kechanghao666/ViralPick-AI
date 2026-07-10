from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
SOURCE = BASE_DIR / "docs" / "requirements_analysis.md"
OUTPUT = BASE_DIR / "docs" / "ViralPick AI 需求分析文档.docx"

FONT_CN = "微软雅黑"
FONT_BODY = "微软雅黑"
COLOR_DARK = "0B1F2A"
COLOR_TEAL = "00D6C9"
COLOR_TEAL_DARK = "064E4A"
COLOR_LIGHT = "EAF7F5"
COLOR_MUTED = "64748B"
COLOR_TABLE_HEAD = "0F766E"
COLOR_TABLE_ALT = "F3FAF9"
COLOR_BORDER = "C7DAD8"


def set_run_font(run, size: int | None = None, bold: bool = False, color: str | None = None) -> None:
    run.font.name = FONT_BODY
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = COLOR_BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_border(paragraph, color: str = COLOR_TEAL, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_BODY)
    r_fonts.set(qn("w:hAnsi"), FONT_BODY)
    r_fonts.set(qn("w:eastAsia"), FONT_CN)
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_mixed_text(paragraph, text: str, size: int = 10.5, bold: bool = False) -> None:
    pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
    position = 0
    for match in pattern.finditer(text):
        before = text[position : match.start()]
        if before:
            run = paragraph.add_run(before)
            set_run_font(run, size=size, bold=bold)
        add_hyperlink(paragraph, match.group(1), match.group(2))
        position = match.end()
    tail = text[position:]
    if tail:
        run = paragraph.add_run(tail)
        set_run_font(run, size=size, bold=bold)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            if row_index == 0:
                set_cell_shading(cell, COLOR_TABLE_HEAD)
            elif row_index % 2 == 0:
                set_cell_shading(cell, COLOR_TABLE_ALT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                run = paragraph.add_run(value.strip())
                set_run_font(run, size=9, bold=row_index == 0, color="FFFFFF" if row_index == 0 else COLOR_DARK)
    doc.add_paragraph()


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E6FFFB")
    set_cell_borders(cell, COLOR_TEAL, "10")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=True, color=COLOR_TEAL_DARK)
    doc.add_paragraph()


def add_flow_box(doc: Document, lines: list[str]) -> None:
    edges: list[str] = []
    for line in lines:
        clean = line.strip()
        if "-->" not in clean:
            continue
        labels = re.findall(r'\["([^"]+)"\]', clean)
        if len(labels) >= 2:
            edges.append(f"{labels[0]} → {labels[1]}")
    if not edges:
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("流程关系")
    set_run_font(run, size=10, bold=True, color=COLOR_TEAL_DARK)
    table = doc.add_table(rows=len(edges), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, edge in enumerate(edges):
        cell = table.cell(index, 0)
        set_cell_borders(cell, COLOR_TEAL, "8")
        set_cell_shading(cell, "F0FDFA" if index % 2 == 0 else "FFFFFF")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(edge)
        set_run_font(run, size=10, bold=True if index == 0 else False, color=COLOR_DARK)
    doc.add_paragraph()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = FONT_BODY
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Heading 1", 18, COLOR_TEAL_DARK),
        ("Heading 2", 14, COLOR_DARK),
        ("Heading 3", 12, COLOR_TEAL_DARK),
    ]:
        style = styles[name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("ViralPick AI 需求分析文档  |  第 ")
    set_run_font(run, size=9, color=COLOR_MUTED)
    add_field(footer, "PAGE")
    run = footer.add_run(" 页")
    set_run_font(run, size=9, color=COLOR_MUTED)
    return doc


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("ViralPick AI")
    set_run_font(run, size=34, bold=True, color=COLOR_TEAL_DARK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("需求分析文档")
    set_run_font(run, size=24, bold=True, color=COLOR_DARK)

    line = doc.add_paragraph()
    set_paragraph_border(line, COLOR_TEAL, "24")

    desc = doc.add_paragraph()
    desc.paragraph_format.space_before = Pt(18)
    desc.paragraph_format.line_spacing = 1.4
    run = desc.add_run("基于 Python 的多平台电商智能经营与增长决策平台")
    set_run_font(run, size=14, bold=True, color=COLOR_DARK)

    meta = [
        ["文档性质", "项目实施前需求分析"],
        ["调研对象", "多平台网店经营者、店铺运营人员、商品运营人员"],
        ["资料检索日期", "2026-07-03"],
        ["生成日期", date.today().isoformat()],
    ]
    doc.add_paragraph()
    add_table(doc, [["项目", "内容"], *meta])
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    note.paragraph_format.left_indent = Cm(0.2)
    set_paragraph_shading(note, "F0FDFA")
    run = note.add_run("本文档聚焦开发前需求分析：先调研行业问题，再推导用户需求、系统需求和验收标准。")
    set_run_font(run, size=11, bold=True, color=COLOR_TEAL_DARK)
    doc.add_page_break()


def add_manual_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    paragraph = doc.add_paragraph("目录", style="Heading 1")
    set_paragraph_border(paragraph)
    for level, text in headings:
        if level != 2:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=COLOR_DARK)
    doc.add_page_break()


def parse_markdown(md: str) -> tuple[list, list[tuple[int, str]]]:
    elements = []
    headings = []
    lines = md.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip() or line.strip() == "---":
            index += 1
            continue
        if line.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].startswith(">"):
                quote_lines.append(lines[index].lstrip("> ").rstrip())
                index += 1
            elements.append(("quote", quote_lines))
            continue
        if line.startswith("```mermaid"):
            flow_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                flow_lines.append(lines[index])
                index += 1
            index += 1
            elements.append(("flow", flow_lines))
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            headings.append((level, text))
            elements.append(("heading", level, text))
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = []
            for raw in table_lines:
                cells = [cell.strip() for cell in raw.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells):
                    continue
                rows.append(cells)
            elements.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while index < len(lines) and lines[index].startswith("- "):
                items.append(lines[index][2:].strip())
                index += 1
            elements.append(("bullets", items))
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[index]).strip())
                index += 1
            elements.append(("numbers", items))
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            current = lines[index].rstrip()
            if not current.strip() or current.startswith(("#", "|", "- ", ">", "```")) or re.match(r"^\d+\.\s+", current):
                break
            paragraph_lines.append(current)
            index += 1
        elements.append(("paragraph", " ".join(part.strip() for part in paragraph_lines)))
    return elements, headings


def render_elements(doc: Document, elements: list) -> None:
    skip_cover_heading = True
    for element in elements:
        kind = element[0]
        if kind == "heading":
            _, level, text = element
            if level == 1 and skip_cover_heading:
                skip_cover_heading = False
                continue
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            paragraph = doc.add_paragraph(text, style=style)
            if level == 2:
                set_paragraph_border(paragraph)
            continue
        if kind == "quote":
            for line in element[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.space_after = Pt(3)
                set_paragraph_shading(p, "F8FAFC")
                add_mixed_text(p, line, size=9.5, bold=True)
            continue
        if kind == "paragraph":
            text = element[1]
            if text.startswith("**需求结论：**"):
                add_callout(doc, text.replace("**", ""))
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0.74)
                p.paragraph_format.line_spacing = 1.35
                p.paragraph_format.space_after = Pt(5)
                add_mixed_text(p, text.replace("**", ""), size=10.5)
            continue
        if kind == "bullets":
            for item in element[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_mixed_text(p, item, size=10)
            continue
        if kind == "numbers":
            for item in element[1]:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(2)
                add_mixed_text(p, item, size=10)
            continue
        if kind == "table":
            add_table(doc, element[1])
            continue
        if kind == "flow":
            add_flow_box(doc, element[1])
            continue


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    elements, headings = parse_markdown(markdown)
    doc = setup_document()
    add_cover(doc)
    add_manual_toc(doc, headings)
    render_elements(doc, elements)
    doc.core_properties.title = "ViralPick AI 需求分析文档"
    doc.core_properties.subject = "多平台电商智能经营与增长决策平台需求分析"
    doc.core_properties.author = "ViralPick AI"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
